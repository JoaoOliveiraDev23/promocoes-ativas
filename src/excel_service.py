"""Leitura, validação e atualização do modelo Excel de promoções ativas.

Princípios:

* O modelo enviado é aberto com openpyxl e **atualizado**, nunca recriado.
* O arquivo original nunca é alterado: tudo é feito em memória (``bytes``) e a
  saída é um conteúdo novo.
* A estrutura do modelo é verificada ANTES de qualquer modificação e o
  resultado é conferido DEPOIS de gerado.
* Texto colado nunca é interpretado como fórmula.

Limitações conhecidas do openpyxl (documentadas no README):

* ``xl/printerSettings/printerSettings1.bin`` (configuração binária da impressora)
  não é preservado; tamanho de papel e orientação da página são mantidos.
* Atributos internos de revisão do Excel (``xr:uid`` etc.) são descartados;
  não têm efeito visual nem funcional.
"""

from __future__ import annotations

import io
import os
import tempfile
import zipfile
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Sequence

from openpyxl import load_workbook
from openpyxl.styles.cell_style import StyleArray
from openpyxl.styles.numbers import is_date_format
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet

from .constants import COLUMN_COUNT, EXPECTED_HEADERS, FIRST_DATA_ROW, OUTPUT_PREFIX
from .models import PromoRecord

_TEXT_PREFIXES_NEEDING_QUOTE = ("=", "+", "-", "@")


# --------------------------------------------------------------------------- #
# Erros
# --------------------------------------------------------------------------- #
class ExcelServiceError(Exception):
    """Erro base; a mensagem é amigável e pode ser mostrada ao usuário."""


class TemplateError(ExcelServiceError):
    """O arquivo enviado não é um modelo compatível."""


class GenerationError(ExcelServiceError):
    """Falha ao gerar ou conferir o arquivo de saída."""


# --------------------------------------------------------------------------- #
# Estruturas
# --------------------------------------------------------------------------- #
@dataclass
class TemplateInfo:
    """Resumo da estrutura encontrada no modelo."""

    sheet_name: str
    current_records: int  # linhas de dados preenchidas hoje
    data_end_row: int  # última linha da área de dados atual (inclui linhas vazias formatadas)
    table_name: str | None = None
    table_ref: str | None = None
    warnings: list[str] = field(default_factory=list)


@dataclass
class _Layout:
    ws: Worksheet
    table: Table | None
    track_autofilter: bool  # autofiltro de planilha (sem tabela) cobrindo A1:F...
    info: TemplateInfo


# --------------------------------------------------------------------------- #
# Nome do arquivo de saída
# --------------------------------------------------------------------------- #
def build_output_filename(when: date | datetime | None = None) -> str:
    """``PROMOCOES-ATIVAS DD-MM-AAAA.xlsx`` com a data local de geração."""
    when = when or datetime.now()
    return f"{OUTPUT_PREFIX} {when:%d-%m-%Y}.xlsx"


# --------------------------------------------------------------------------- #
# Abertura e análise do modelo
# --------------------------------------------------------------------------- #
def _norm(value: object) -> str:
    return " ".join(str(value).strip().upper().split()) if value is not None else ""


def _open_workbook(data: bytes):
    if not data:
        raise TemplateError("O arquivo enviado está vazio.")
    buffer = io.BytesIO(data)
    if not zipfile.is_zipfile(buffer):
        raise TemplateError(
            "O arquivo não é um .xlsx válido (pode estar corrompido, protegido por senha "
            "ou ser de outro formato, como .xls)."
        )
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = set(zf.namelist())
    if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
        raise TemplateError("O arquivo não tem a estrutura de uma pasta de trabalho do Excel (.xlsx).")
    if "xl/vbaProject.bin" in names:
        raise TemplateError("Arquivos com macros (.xlsm) não são aceitos. Envie o modelo em .xlsx.")
    try:
        return load_workbook(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001 - arquivo externo: qualquer falha vira mensagem clara
        raise TemplateError(f"Não foi possível ler o arquivo Excel ({type(exc).__name__}).") from exc


def _headers_match(ws: Worksheet) -> bool:
    return all(_norm(ws.cell(1, c + 1).value) == EXPECTED_HEADERS[c] for c in range(COLUMN_COUNT))


def _has_value(value: object) -> bool:
    return value is not None and value != ""


def _analyze(wb) -> _Layout:
    # --- aba de dados --------------------------------------------------------
    matches = [ws for ws in wb.worksheets if _headers_match(ws)]
    if not matches:
        first = wb.worksheets[0] if wb.worksheets else None
        found = [str(first.cell(1, c + 1).value or "") for c in range(COLUMN_COUNT)] if first else []
        raise TemplateError(
            "Nenhuma aba do arquivo tem os cabeçalhos esperados na linha 1 "
            f"({', '.join(EXPECTED_HEADERS)}). "
            f"Cabeçalhos encontrados na primeira aba: {', '.join(found) or '(vazio)'}."
        )
    if len(matches) > 1:
        names = ", ".join(f"'{m.title}'" for m in matches)
        raise TemplateError(f"Mais de uma aba tem os cabeçalhos esperados ({names}); não é possível escolher com segurança.")
    ws = matches[0]
    warnings: list[str] = []

    if ws.protection.sheet:
        warnings.append("A aba está protegida; a proteção será mantida no arquivo de saída.")

    # --- tabela estruturada --------------------------------------------------
    table: Table | None = None
    for t in ws.tables.values():
        min_col, min_row, max_col, max_row = range_boundaries(t.ref)
        if min_col == 1 and min_row == 1 and max_col == COLUMN_COUNT:
            table = t
        elif min_col <= COLUMN_COUNT and max_col >= 1:
            raise TemplateError(
                f"A aba tem outra tabela ('{t.displayName}', {t.ref}) sobre as colunas A:F; "
                "não é possível atualizar com segurança."
            )
    if table is not None:
        if table.totalsRowCount:
            raise TemplateError(f"A tabela '{table.displayName}' tem linha de totais, o que não é suportado.")
        cols = [_norm(c.name) for c in table.tableColumns]
        if cols != list(EXPECTED_HEADERS):
            raise TemplateError(
                f"As colunas da tabela '{table.displayName}' não correspondem aos cabeçalhos esperados."
            )

    # --- autofiltro de planilha (quando não há tabela) -------------------------
    track_autofilter = False
    filter_end = 1
    if table is None and ws.auto_filter.ref:
        min_col, min_row, max_col, max_row = range_boundaries(ws.auto_filter.ref)
        if min_col == 1 and min_row == 1 and max_col == COLUMN_COUNT:
            track_autofilter = True
            filter_end = max_row

    # --- extensão atual da área de dados --------------------------------------
    last_value_row = 1
    for r, row in enumerate(ws.iter_rows(min_row=2, max_col=COLUMN_COUNT, values_only=True), start=2):
        if any(_has_value(v) for v in row):
            last_value_row = r
    candidates = [last_value_row, filter_end]
    if table is not None:
        candidates.append(range_boundaries(table.ref)[3])
    data_end = max(candidates)

    # --- verificações bloqueantes e avisos na área de dados ----------------------
    for rng in ws.merged_cells.ranges:
        if rng.max_row >= FIRST_DATA_ROW and rng.min_col <= COLUMN_COUNT:
            raise TemplateError(f"Há células mescladas na área de dados ({rng.coord}); não é possível atualizar com segurança.")

    records = 0
    has_comment_or_link = False
    for row in ws.iter_rows(min_row=FIRST_DATA_ROW, max_row=max(data_end, FIRST_DATA_ROW), max_col=COLUMN_COUNT):
        filled = False
        for cell in row:
            if cell.data_type == "f":
                raise TemplateError(
                    f"A célula {cell.coordinate} contém fórmula na área de dados. "
                    "A aplicação não substitui fórmulas; remova-a do modelo ou use outro arquivo."
                )
            if _has_value(cell.value):
                filled = True
            if cell.comment is not None or cell.hyperlink is not None:
                has_comment_or_link = True
        records += filled

    if has_comment_or_link:
        warnings.append("Há comentários ou links em células da área de dados; eles podem permanecer em linhas que mudarem de conteúdo.")

    outside = any(
        _has_value(cell.value)
        for row in ws.iter_rows(min_col=COLUMN_COUNT + 1)
        for cell in row
    )
    if outside:
        warnings.append("Há conteúdo fora das colunas A:F; ele não será alterado.")
    if len(ws.conditional_formatting):
        warnings.append("O modelo tem formatação condicional; seus intervalos não são ajustados.")
    if ws.data_validations is not None and ws.data_validations.dataValidation:
        warnings.append("O modelo tem validação de dados; seus intervalos não são ajustados.")
    if ws.print_area or len(wb.defined_names) or len(ws.defined_names):
        warnings.append("O modelo tem área de impressão ou nomes definidos; seus intervalos não são ajustados.")

    info = TemplateInfo(
        sheet_name=ws.title,
        current_records=records,
        data_end_row=data_end,
        table_name=table.displayName if table is not None else None,
        table_ref=table.ref if table is not None else None,
        warnings=warnings,
    )
    return _Layout(ws, table, track_autofilter, info)


def inspect_template(data: bytes) -> TemplateInfo:
    """Valida o arquivo enviado (sem alterá-lo) e descreve sua estrutura."""
    return _analyze(_open_workbook(data)).info


# --------------------------------------------------------------------------- #
# Atualização
# --------------------------------------------------------------------------- #
def _set_text(cell, text: str) -> None:
    if text == "":
        cell.value = None
        return
    cell.value = text
    cell.data_type = "s"  # nunca interpretar como fórmula, mesmo que comece com "="
    if text.startswith(_TEXT_PREFIXES_NEEDING_QUOTE):
        cell.quotePrefix = True  # se alguém editar a célula, o Excel continua tratando como texto


def _set_date(cell, value: date) -> None:
    cell.value = datetime(value.year, value.month, value.day)
    if not is_date_format(cell.number_format):
        cell.number_format = "dd/mm/yyyy"


def _style_signature(cell) -> tuple:
    border = cell.border
    return (
        cell.number_format,
        cell.alignment.horizontal,
        cell.alignment.vertical,
        cell.alignment.wrap_text,
        border.left.style,
        border.right.style,
        border.top.style,
        border.bottom.style,
        cell.font.name,
        cell.font.sz,
        cell.font.b,
        cell.font.i,
        cell.font.u,
        cell.fill.fill_type,
        cell.fill.fgColor.rgb if cell.fill.fill_type else None,
    )


def update_workbook(template: bytes, records: Sequence[PromoRecord], *, shrink_table: bool = True) -> bytes:
    """Gera o conteúdo do arquivo atualizado a partir do modelo (em ``bytes``).

    ``shrink_table=True`` ajusta a tabela estruturada exatamente ao número de
    registros novos (as linhas vazias que sobravam abaixo deixam de fazer parte
    da tabela e perdem a borda). Com ``False`` a tabela nunca encolhe.
    """
    if not records:
        raise GenerationError("Não há registros para gravar.")

    wb = _open_workbook(template)
    layout = _analyze(wb)
    ws, table = layout.ws, layout.table
    old_end = layout.info.data_end_row
    count = len(records)
    last_new_row = FIRST_DATA_ROW + count - 1

    # Estilo-padrão: o da primeira linha de dados do modelo (linha 2).
    row_styles = [copy(ws.cell(FIRST_DATA_ROW, c)._style) for c in range(1, COLUMN_COUNT + 1)]
    row_height = ws.row_dimensions[FIRST_DATA_ROW].height

    header_before = [_style_signature(ws.cell(1, c)) for c in range(1, COLUMN_COUNT + 1)]
    widths_before = {k: v.width for k, v in ws.column_dimensions.items()}
    freeze_before = ws.freeze_panes
    sheets_before = wb.sheetnames

    # 1) Limpa os valores antigos (mantém estilos).
    for r in range(FIRST_DATA_ROW, old_end + 1):
        for c in range(1, COLUMN_COUNT + 1):
            ws.cell(r, c).value = None

    # 2) Grava os novos registros com a formatação da linha-modelo.
    for offset, rec in enumerate(records):
        r = FIRST_DATA_ROW + offset
        cells = [ws.cell(r, c) for c in range(1, COLUMN_COUNT + 1)]
        for cell, style in zip(cells, row_styles):
            cell._style = copy(style)
        cells[0].value = rec.cod
        _set_text(cells[1], rec.nome)
        _set_text(cells[2], rec.detalhes)
        _set_date(cells[3], rec.data_inicial)
        _set_date(cells[4], rec.data_final)
        cells[5].value = rec.prioridade
        if row_height is not None and r > old_end:
            ws.row_dimensions[r].height = row_height

    # 3) Linhas antigas que sobraram abaixo dos novos dados.
    new_end = last_new_row
    if old_end > last_new_row:
        if shrink_table:
            for r in range(last_new_row + 1, old_end + 1):
                for c in range(1, COLUMN_COUNT + 1):
                    ws._cells.pop((r, c), None)  # remove a célula (valor e formatação)
        else:
            new_end = old_end

    # 4) Intervalos da tabela / autofiltro.
    new_ref = f"A1:F{new_end}"
    if table is not None:
        table.ref = new_ref
        if table.autoFilter is not None:
            table.autoFilter.ref = new_ref
    elif layout.track_autofilter:
        ws.auto_filter.ref = new_ref

    # 5) Visual: a janela do modelo abria rolada (A73) com seleção antiga (A2:F99);
    #    com outra quantidade de linhas isso poderia abrir numa área vazia.
    view = ws.sheet_view
    if view.pane is None:
        view.topLeftCell = "A1"
        for selection in view.selection:
            selection.activeCell = "A1"
            selection.sqref = "A1"

    buffer = io.BytesIO()
    try:
        wb.save(buffer)
    except Exception as exc:  # noqa: BLE001
        raise GenerationError(f"Falha ao gerar o arquivo Excel ({type(exc).__name__}).") from exc
    output = buffer.getvalue()

    _verify_output(
        output,
        records,
        sheet_name=ws.title,
        sheets_before=sheets_before,
        table_name=table.displayName if table is not None else None,
        table_style=table.tableStyleInfo.name if table is not None and table.tableStyleInfo else None,
        expected_ref=new_ref if (table is not None or layout.track_autofilter) else None,
        has_table=table is not None,
        old_end=old_end,
        new_end_rows=last_new_row,
        header_before=header_before,
        row_signature=[_style_signature(ws.cell(FIRST_DATA_ROW, c)) for c in range(1, COLUMN_COUNT + 1)],
        widths_before=widths_before,
        freeze_before=freeze_before,
    )
    return output


# --------------------------------------------------------------------------- #
# Conferência pós-geração
# --------------------------------------------------------------------------- #
def _verify_output(
    output: bytes,
    records: Sequence[PromoRecord],
    *,
    sheet_name: str,
    sheets_before: list[str],
    table_name: str | None,
    table_style: str | None,
    expected_ref: str | None,
    has_table: bool,
    old_end: int,
    new_end_rows: int,
    header_before: list[tuple],
    row_signature: list[tuple],
    widths_before: dict,
    freeze_before: str | None,
) -> None:
    def fail(reason: str):
        raise GenerationError(f"A conferência do arquivo gerado falhou: {reason} O arquivo não foi disponibilizado.")

    try:
        wb = load_workbook(io.BytesIO(output))
    except Exception as exc:  # noqa: BLE001
        raise GenerationError("O arquivo gerado não pôde ser reaberto para conferência.") from exc

    if wb.sheetnames != sheets_before:
        fail("os nomes das abas mudaram.")
    ws = wb[sheet_name]

    if [_norm(ws.cell(1, c).value) for c in range(1, COLUMN_COUNT + 1)] != list(EXPECTED_HEADERS):
        fail("os cabeçalhos foram alterados.")
    if [_style_signature(ws.cell(1, c)) for c in range(1, COLUMN_COUNT + 1)] != header_before:
        fail("a formatação dos cabeçalhos foi alterada.")
    if {k: v.width for k, v in ws.column_dimensions.items()} != widths_before:
        fail("as larguras das colunas foram alteradas.")
    if ws.freeze_panes != freeze_before:
        fail("o congelamento de painéis foi alterado.")

    for offset, rec in enumerate(records):
        r = FIRST_DATA_ROW + offset
        got = [ws.cell(r, c).value for c in range(1, COLUMN_COUNT + 1)]
        want = [
            rec.cod,
            rec.nome,
            rec.detalhes or None,
            datetime(rec.data_inicial.year, rec.data_inicial.month, rec.data_inicial.day),
            datetime(rec.data_final.year, rec.data_final.month, rec.data_final.day),
            rec.prioridade,
        ]
        if got != want:
            fail(f"a linha {r} não contém os valores esperados.")
        if [_style_signature(ws.cell(r, c)) for c in range(1, COLUMN_COUNT + 1)] != row_signature:
            fail(f"a formatação da linha {r} difere do padrão do modelo.")

    for r in range(new_end_rows + 1, max(old_end, new_end_rows) + 1):
        if any(_has_value(ws.cell(r, c).value) for c in range(1, COLUMN_COUNT + 1)):
            fail(f"sobraram valores antigos na linha {r}.")

    if has_table:
        table = ws.tables.get(table_name)
        if table is None:
            fail("a tabela estruturada desapareceu.")
        if table.ref != expected_ref or (table.autoFilter is not None and table.autoFilter.ref != expected_ref):
            fail("o intervalo da tabela estruturada está incorreto.")
        if table.tableStyleInfo is None or table.tableStyleInfo.name != table_style:
            fail("o estilo da tabela estruturada foi alterado.")
    elif expected_ref is not None and ws.auto_filter.ref != expected_ref:
        fail("o intervalo do filtro automático está incorreto.")


# --------------------------------------------------------------------------- #
# Versão baseada em arquivos (útil para scripts e testes)
# --------------------------------------------------------------------------- #
def update_workbook_file(
    source: str | os.PathLike,
    records: Sequence[PromoRecord],
    output_dir: str | os.PathLike,
    *,
    when: date | datetime | None = None,
    shrink_table: bool = True,
    overwrite: bool = False,
) -> Path:
    """Lê ``source`` (somente leitura) e grava o arquivo atualizado em ``output_dir``.

    O original nunca é alterado: se o destino coincidir com ele, a operação é
    recusada. Um arquivo de saída já existente só é substituído com ``overwrite=True``.
    """
    source_path = Path(source).resolve()
    try:
        data = source_path.read_bytes()
    except OSError as exc:
        raise TemplateError(f"Não foi possível ler o arquivo: {exc.strerror or type(exc).__name__}.") from exc

    output = update_workbook(data, records, shrink_table=shrink_table)

    destination = (Path(output_dir) / build_output_filename(when)).resolve()
    if destination == source_path:
        raise GenerationError("O arquivo de saída teria o mesmo caminho do original, que nunca é sobrescrito.")
    if destination.exists() and not overwrite:
        raise GenerationError(f"Já existe um arquivo chamado '{destination.name}' em {destination.parent}.")

    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
        fd, tmp_name = tempfile.mkstemp(suffix=".tmp", dir=destination.parent)
        try:
            with os.fdopen(fd, "wb") as handle:
                handle.write(output)
            os.replace(tmp_name, destination)
        finally:
            if os.path.exists(tmp_name):
                os.remove(tmp_name)
    except OSError as exc:
        raise GenerationError(f"Não foi possível gravar o arquivo de saída: {exc.strerror or type(exc).__name__}.") from exc
    return destination
