"""Fluxo único texto -> planilha, usando o modelo interno (assets/)."""

from __future__ import annotations

import hashlib
import io
from datetime import date

import pytest
from openpyxl import load_workbook

from src.excel_service import ExcelServiceError, TemplateError, inspect_template
from src.generator import (
    DEFAULT_TEMPLATE_PATH,
    MAX_ERRORS_SHOWN,
    InvalidInput,
    generate_workbook,
    load_template,
)

LINHA = "{cod}\tPromoção {cod} – R$ 2,87\tDetalhe {cod}\t2026-03-10\t2026-12-31\t{prio}"


def texto(n: int) -> str:
    return "\n".join(LINHA.format(cod=500 + i, prio=1 + i % 4) for i in range(n)) + "\n"


def sha(path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_modelo_interno_existe_e_e_valido():
    assert DEFAULT_TEMPLATE_PATH.is_file()
    info = inspect_template(load_template())
    assert info.sheet_name == "Planilha1"
    assert info.table_name == "Tabela1"
    assert info.current_records == 98


def test_gera_planilha_com_nome_padrao_e_data_local():
    gerado = generate_workbook(texto(3), now=date(2026, 10, 9))
    assert gerado.filename == "PROMOCOES-ATIVAS 09-10-2026.xlsx"
    assert gerado.record_count == 3


def test_substitui_integralmente_os_registros_do_modelo():
    ws = load_workbook(io.BytesIO(generate_workbook(texto(5)).data)).active
    assert [ws.cell(r, 1).value for r in range(2, 8)] == [500, 501, 502, 503, 504, None]
    assert ws.tables["Tabela1"].ref == "A1:F6"
    # nada dos 98 registros antigos sobrou (nem valores nem formatação) nas colunas A:F
    for linha in ws.iter_rows(min_row=7, max_row=130, max_col=6):
        for celula in linha:
            assert celula.value is None and not celula.has_style, celula.coordinate


def test_ajusta_linhas_quando_ha_mais_registros_que_o_modelo():
    ws = load_workbook(io.BytesIO(generate_workbook(texto(130)).data)).active
    assert ws.cell(131, 1).value == 629 and ws.cell(132, 1).value is None
    assert ws.tables["Tabela1"].ref == "A1:F131"
    assert ws.cell(131, 4).number_format == ws.cell(2, 4).number_format
    assert ws.cell(131, 1).border.left.style == ws.cell(2, 1).border.left.style


def test_modelo_interno_nunca_e_alterado(real_model_path):
    antes = sha(real_model_path)
    generate_workbook(texto(10))
    with pytest.raises(InvalidInput):
        generate_workbook("lixo")
    assert sha(real_model_path) == antes


def test_modelo_ausente_gera_mensagem_clara(tmp_path):
    with pytest.raises(TemplateError, match="modelo interno não foi encontrado") as exc:
        generate_workbook(texto(2), template_path=tmp_path / "nao_existe.xlsx")
    assert "assets/PROMOCOES-ATIVAS-modelo.xlsx" in str(exc.value)


def test_modelo_interno_invalido_gera_mensagem_clara(tmp_path):
    falso = tmp_path / "modelo.xlsx"
    falso.write_bytes(b"nao sou um excel")
    with pytest.raises(TemplateError, match="modelo interno"):
        generate_workbook(texto(2), template_path=falso)


@pytest.mark.parametrize("vazio", ["", "   ", "\n\n\t\n"])
def test_texto_vazio_pede_para_colar(vazio):
    with pytest.raises(InvalidInput, match="Cole o resultado"):
        generate_workbook(vazio)


def test_validacao_automatica_informa_linha_e_motivo():
    with pytest.raises(InvalidInput) as exc:
        generate_workbook(texto(2) + "abc\tNome\t\t10/03/2025\t2026-12-31\t1")
    linhas = exc.value.lines
    assert linhas[0].startswith("Linha 3 [COD]") or linhas[0].startswith("Linha 3 [DATA INICIAL]")
    assert any("dia/mês/ano" in l for l in linhas)


def test_cabecalho_colado_impede_a_geracao():
    cab = "COD\tNOME\tDETALHES\tDATA INICIAL\tDATA FINAL\tPRIORIDADE\n"
    with pytest.raises(InvalidInput, match="cabeçalhos"):
        generate_workbook(cab + texto(2))


def test_limita_a_quantidade_de_erros_mostrados():
    ruim = "\n".join("x\ty" for _ in range(20))
    with pytest.raises(InvalidInput) as exc:
        generate_workbook(ruim)
    assert len(exc.value.lines) == MAX_ERRORS_SHOWN and exc.value.total == 20


def test_avisos_nao_impedem_a_geracao():
    gerado = generate_workbook("1\tA\tx\t2026-05-01\t2026-04-01\t1\n1\tB\t\t2026-01-01\t2026-02-01\t1")
    assert gerado.record_count == 2 and gerado.warnings


def test_todos_os_erros_sao_excecoes_amigaveis():
    for entrada in ("", "a\tb", "1\tA\tx\t2026-02-30\t2026-12-31\t1"):
        with pytest.raises(ExcelServiceError):
            generate_workbook(entrada)
