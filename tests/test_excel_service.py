from __future__ import annotations

import hashlib
import io
from datetime import date, datetime

import pytest
from openpyxl import load_workbook

from src.data_parser import parse_promotions
from src.excel_service import (
    GenerationError,
    TemplateError,
    _style_signature,
    build_output_filename,
    inspect_template,
    update_workbook,
    update_workbook_file,
)
from src.models import PromoRecord
from tests.conftest import HEADERS, WIDTHS, make_records


def sha(path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def reopen(data: bytes):
    return load_workbook(io.BytesIO(data)).active


def values(ws, row: int) -> list:
    return [ws.cell(row, c).value for c in range(1, 7)]


# ----------------------------- nome do arquivo ----------------------------- #
def test_nome_do_arquivo_segue_o_padrao():
    assert build_output_filename(date(2026, 10, 9)) == "PROMOCOES-ATIVAS 09-10-2026.xlsx"
    assert build_output_filename(datetime(2027, 1, 2, 23, 59)) == "PROMOCOES-ATIVAS 02-01-2027.xlsx"


def test_nome_do_arquivo_sem_argumento_usa_a_data_local_de_hoje():
    assert build_output_filename() == f"PROMOCOES-ATIVAS {datetime.now():%d-%m-%Y}.xlsx"


# ------------------------------ inspeção ------------------------------------ #
def test_inspecao_descreve_a_estrutura(model_path):
    info = inspect_template(model_path.read_bytes())
    assert info.sheet_name == "Planilha1"
    assert info.current_records == 5
    assert info.table_name == "Tabela1" and info.table_ref == "A1:F21"
    assert info.data_end_row == 21
    assert info.warnings == []


def test_inspecao_nao_altera_o_arquivo(model_path):
    before = sha(model_path)
    inspect_template(model_path.read_bytes())
    assert sha(model_path) == before


# -------------------------- substituição dos dados -------------------------- #
def test_substitui_integralmente_os_registros_antigos(model_path):
    novos = make_records(5)
    ws = reopen(update_workbook(model_path.read_bytes(), novos))
    for i, rec in enumerate(novos):
        assert values(ws, 2 + i) == [
            rec.cod,
            rec.nome,
            rec.detalhes,
            datetime(2026, 1, 1 + i % 28),
            datetime(2026, 12, 31),
            rec.prioridade,
        ]
    texto = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "Antiga" not in texto and "Detalhe antigo" not in texto


def test_remove_registros_excedentes_antigos(model_factory):
    path = model_factory(records=15, table_rows=15)
    ws = reopen(update_workbook(path.read_bytes(), make_records(3)))
    assert values(ws, 4)[0] == 1002
    for r in range(5, 17):
        assert values(ws, r) == [None] * 6
    assert ws.tables["Tabela1"].ref == "A1:F4"


def test_linhas_vazias_formatadas_que_sobram_abaixo_da_tabela_sao_removidas(model_path):
    ws = reopen(update_workbook(model_path.read_bytes(), make_records(3)))
    assert ws.max_row == 4  # antes de acessar células (ws.cell cria células ao ler)
    # o modelo tinha linhas vazias com borda até a 21; abaixo dos 3 registros não pode sobrar nada
    for r in range(5, 22):
        for c in range(1, 7):
            assert ws.cell(r, c).value is None
            assert not ws.cell(r, c).has_style, f"célula {ws.cell(r, c).coordinate} ainda tem formatação"


def test_inclui_novas_linhas_com_a_formatacao_do_padrao(model_path):
    ws = reopen(update_workbook(model_path.read_bytes(), make_records(40)))
    assert values(ws, 41)[0] == 1039 and values(ws, 42) == [None] * 6
    assert ws.tables["Tabela1"].ref == "A1:F41"
    modelo = _style_signature(ws.cell(2, 1)), _style_signature(ws.cell(2, 4))
    for r in (21, 22, 30, 41):  # 21 era a última linha da tabela; 22+ são novas
        assert (_style_signature(ws.cell(r, 1)), _style_signature(ws.cell(r, 4))) == modelo
    assert ws.cell(41, 4).number_format == "mm-dd-yy"
    assert ws.cell(41, 2).alignment.horizontal == "left"
    assert ws.cell(41, 1).border.left.style == "thin"


def test_formata_e_alinha_cada_coluna_como_no_modelo(model_path):
    antes = load_workbook(model_path).active
    ws = reopen(update_workbook(model_path.read_bytes(), make_records(3)))
    for r in (2, 3, 4):
        for c in range(1, 7):
            assert _style_signature(ws.cell(r, c)) == _style_signature(antes.cell(2, c))


def test_datas_sao_gravadas_como_data_real(model_path):
    ws = reopen(update_workbook(model_path.read_bytes(), make_records(1)))
    assert isinstance(ws["D2"].value, datetime) and ws["D2"].is_date
    assert ws["D2"].number_format == "mm-dd-yy"


def test_cod_e_prioridade_sao_numeros_e_textos_continuam_textos(model_path):
    ws = reopen(update_workbook(model_path.read_bytes(), make_records(1)))
    assert isinstance(ws["A2"].value, int) and isinstance(ws["F2"].value, int)
    assert isinstance(ws["B2"].value, str) and isinstance(ws["C2"].value, str)


def test_detalhes_vazio_vira_celula_vazia_com_formatacao(model_path):
    rec = PromoRecord(1, "Nome", "", date(2026, 1, 1), date(2026, 1, 2), 1)
    ws = reopen(update_workbook(model_path.read_bytes(), [rec]))
    assert ws["C2"].value is None and ws["C2"].border.left.style == "thin"


# ------------------------- segurança do conteúdo ----------------------------- #
def test_texto_que_comeca_com_igual_nao_vira_formula(model_path):
    result = parse_promotions('1\t=1+1\t=HYPERLINK("http://x","y")\t2026-01-01\t2026-01-02\t1')
    data = update_workbook(model_path.read_bytes(), result.records)
    ws = reopen(data)
    assert ws["B2"].value == "=1+1" and ws["B2"].data_type == "s"
    assert ws["C2"].data_type == "s"
    assert ws["B2"].quotePrefix is True
    # nenhuma fórmula gravada no XML
    import zipfile

    xml = zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml").decode()
    assert "<f>" not in xml and "<f " not in xml


# ------------------------------ preservação ---------------------------------- #
def test_preserva_cabecalhos_formatacao_larguras_e_nome_da_aba(model_factory):
    path = model_factory(records=5, table_rows=20, freeze="A2")
    antes = load_workbook(path).active
    ws = reopen(update_workbook(path.read_bytes(), make_records(8)))
    assert ws.title == "Planilha1"
    assert [ws.cell(1, c).value for c in range(1, 7)] == HEADERS
    for c in range(1, 7):
        assert _style_signature(ws.cell(1, c)) == _style_signature(antes.cell(1, c))
    assert {k: round(v.width, 2) for k, v in ws.column_dimensions.items()} == WIDTHS
    assert ws.freeze_panes == "A2"


def test_tabela_estruturada_atualiza_intervalo_e_preserva_configuracao(model_path):
    ws = reopen(update_workbook(model_path.read_bytes(), make_records(9)))
    tabela = ws.tables["Tabela1"]
    assert tabela.ref == "A1:F10" and tabela.autoFilter.ref == "A1:F10"
    assert tabela.tableStyleInfo.name == "TableStyleMedium16"
    assert tabela.tableStyleInfo.showRowStripes is True
    assert [c.name for c in tabela.tableColumns] == HEADERS


def test_tabela_nao_encolhe_com_shrink_table_false(model_path):
    data = update_workbook(model_path.read_bytes(), make_records(3), shrink_table=False)
    ws = reopen(data)
    assert ws.tables["Tabela1"].ref == "A1:F21"
    assert ws.cell(10, 1).border.left.style == "thin"  # linha vazia continua formatada


def test_modelo_sem_tabela_com_autofiltro_atualiza_o_filtro(model_factory):
    path = model_factory(records=5, table_rows=10, with_table=False)
    wb = load_workbook(path)
    wb.active.auto_filter.ref = "A1:F11"
    wb.save(path)
    ws = reopen(update_workbook(path.read_bytes(), make_records(4)))
    assert ws.auto_filter.ref == "A1:F5"


def test_modelo_sem_tabela_funciona(model_factory):
    path = model_factory(records=5, table_rows=5, with_table=False)
    ws = reopen(update_workbook(path.read_bytes(), make_records(7)))
    assert values(ws, 8)[0] == 1006 and not ws.tables


def test_visualizacao_volta_ao_topo(model_path):
    wb = load_workbook(model_path)
    wb.active.sheet_view.topLeftCell = "A73"
    wb.active.sheet_view.selection[0].sqref = "A2:F99"
    wb.save(model_path)
    ws = reopen(update_workbook(model_path.read_bytes(), make_records(3)))
    assert ws.sheet_view.topLeftCell == "A1"
    assert ws.sheet_view.selection[0].sqref == "A1"


# ------------------------------ arquivo original ------------------------------ #
def test_arquivo_original_permanece_inalterado(model_path, tmp_path):
    antes = sha(model_path)
    saida = tmp_path / "saida"
    destino = update_workbook_file(model_path, make_records(30), saida, when=date(2026, 10, 9))
    assert sha(model_path) == antes
    assert destino.name == "PROMOCOES-ATIVAS 09-10-2026.xlsx" and destino.parent == saida
    assert reopen(destino.read_bytes())["A2"].value == 1000
    assert load_workbook(model_path).active["A2"].value == 9000


def test_recusa_sobrescrever_o_original(model_factory, tmp_path):
    nome = build_output_filename(date(2026, 10, 9))
    original = model_factory(records=5)
    copia = original.parent / nome
    copia.write_bytes(original.read_bytes())
    antes = sha(copia)
    with pytest.raises(GenerationError, match="nunca é sobrescrito"):
        update_workbook_file(copia, make_records(2), copia.parent, when=date(2026, 10, 9), overwrite=True)
    assert sha(copia) == antes


def test_nao_substitui_arquivo_existente_sem_pedir(model_path, tmp_path):
    update_workbook_file(model_path, make_records(2), tmp_path / "o", when=date(2026, 10, 9))
    with pytest.raises(GenerationError, match="Já existe"):
        update_workbook_file(model_path, make_records(2), tmp_path / "o", when=date(2026, 10, 9))


def test_update_workbook_nao_modifica_os_bytes_recebidos(model_path):
    data = model_path.read_bytes()
    copia = bytes(data)
    update_workbook(data, make_records(3))
    assert data == copia


# ------------------------------ arquivos inválidos ------------------------------ #
def test_rejeita_arquivo_que_nao_e_xlsx():
    for lixo in (b"", b"isto nao e um xlsx", b"PK\x03\x04quebrado"):
        with pytest.raises(TemplateError):
            inspect_template(lixo)


def test_rejeita_zip_que_nao_e_excel():
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("qualquer.txt", "x")
    with pytest.raises(TemplateError, match="estrutura"):
        inspect_template(buf.getvalue())


def test_rejeita_cabecalhos_inesperados(model_factory):
    path = model_factory(headers=["CODIGO", "NOME", "DETALHES", "INICIO", "FIM", "PRIORIDADE"])
    with pytest.raises(TemplateError, match="cabeçalhos esperados") as exc:
        inspect_template(path.read_bytes())
    assert "CODIGO" in str(exc.value)


def test_rejeita_formula_na_area_de_dados(model_factory):
    path = model_factory(formula_in_data=True)
    with pytest.raises(TemplateError, match="fórmula"):
        update_workbook(path.read_bytes(), make_records(2))


def test_rejeita_celulas_mescladas_na_area_de_dados(model_path):
    wb = load_workbook(model_path)
    wb.active.merge_cells("B3:C3")
    wb.save(model_path)
    with pytest.raises(TemplateError, match="mescladas"):
        inspect_template(model_path.read_bytes())


def test_rejeita_outra_tabela_sobre_as_colunas(model_factory):
    from openpyxl.worksheet.table import Table

    path = model_factory(records=3, table_rows=3)
    wb = load_workbook(path)
    ws = wb.active
    for col, name in zip("ABC", ["X", "Y", "Z"]):
        ws[f"{col}30"] = name
    ws.add_table(Table(displayName="Outra", ref="A30:C31"))
    wb.save(path)
    with pytest.raises(TemplateError, match="outra tabela"):
        inspect_template(path.read_bytes())


def test_sem_registros_nao_gera(model_path):
    with pytest.raises(GenerationError, match="Não há registros"):
        update_workbook(model_path.read_bytes(), [])


def test_avisa_sobre_conteudo_fora_das_colunas_a_f(model_path):
    wb = load_workbook(model_path)
    wb.active["H5"] = "observação"
    wb.save(model_path)
    info = inspect_template(model_path.read_bytes())
    assert any("fora das colunas" in w for w in info.warnings)


def test_conteudo_fora_das_colunas_a_f_nao_e_alterado(model_path):
    wb = load_workbook(model_path)
    wb.active["H5"] = "observação"
    wb.save(model_path)
    ws = reopen(update_workbook(model_path.read_bytes(), make_records(2)))
    assert ws["H5"].value == "observação"


def test_outras_abas_sao_preservadas(model_path):
    wb = load_workbook(model_path)
    wb.create_sheet("Notas")["A1"] = "nota"
    wb.save(model_path)
    saida = load_workbook(io.BytesIO(update_workbook(model_path.read_bytes(), make_records(2))))
    assert saida.sheetnames == ["Planilha1", "Notas"] and saida["Notas"]["A1"].value == "nota"


# ----------------------- fluxo completo: texto -> planilha ---------------------- #
def test_fluxo_completo_do_texto_ate_a_planilha(model_path):
    texto = (
        "187\tPetit Deli Ciclo 3 - 2025\tGranéis – até R$ 2,87\t2025-03-10\t2026-12-31\t4\n"
        "233\tKIT VITALÍCIO\t\t2025-05-05\t2026-12-31\t1\n"
    )
    result = parse_promotions(texto)
    assert result.is_valid
    ws = reopen(update_workbook(model_path.read_bytes(), result.records))
    assert values(ws, 2) == [187, "Petit Deli Ciclo 3 - 2025", "Granéis – até R$ 2,87", datetime(2025, 3, 10), datetime(2026, 12, 31), 4]
    assert values(ws, 3)[2] is None and values(ws, 4) == [None] * 6


# ------------------------------ modelo real embutido (assets/) ------------------ #
def test_modelo_real_roundtrip_identico(real_model_path):
    """Regrava os mesmos dados do modelo embutido e exige resultado idêntico (valores e estilos)."""
    original = load_workbook(real_model_path).active
    linhas = []
    for r in original.iter_rows(min_row=2, max_col=6, values_only=True):
        if any(v is not None for v in r):
            linhas.append(
                "\t".join([str(r[0]), r[1], r[2] or "", r[3].strftime("%Y-%m-%d"), r[4].strftime("%Y-%m-%d"), str(r[5])])
            )
    antes = sha(real_model_path)
    result = parse_promotions("\n".join(linhas) + "\n")
    assert result.is_valid
    ws = reopen(update_workbook(real_model_path.read_bytes(), result.records))
    n = len(result.records)
    for r in range(1, n + 2):
        for c in range(1, 7):
            assert ws.cell(r, c).value == original.cell(r, c).value
            assert _style_signature(ws.cell(r, c)) == _style_signature(original.cell(r, c))
    assert ws.tables["Tabela1"].ref == f"A1:F{n + 1}"
    assert sha(real_model_path) == antes
