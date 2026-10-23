"""Fixtures compartilhadas.

Os testes NUNCA usam o arquivo real do usuário por padrão: constroem um modelo
sintético com a mesma estrutura do real (aba "Planilha1", tabela "Tabela1" com
estilo TableStyleMedium16, cabeçalhos com borda, colunas centralizadas, datas
no formato de data do Excel, larguras personalizadas e linhas vazias formatadas
sobrando no fim da tabela).

Há também testes contra o modelo real embutido em ``assets/`` (somente leitura).
"""

from __future__ import annotations

import sys
from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.models import PromoRecord  # noqa: E402

THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WIDTHS = {"A": 7.57, "B": 51.57, "C": 72.71, "D": 15.14, "E": 13.57, "F": 15.0}
HEADERS = ["COD", "NOME", "DETALHES", "DATA INICIAL", "DATA FINAL", "PRIORIDADE"]


def make_records(n: int, start_cod: int = 1000) -> list[PromoRecord]:
    return [
        PromoRecord(
            cod=start_cod + i,
            nome=f"Promoção {i} – Ciclo {i % 13}",
            detalhes=f"Compre {i + 2} unidades por R$ {i + 1},99",
            data_inicial=date(2026, 1, 1 + i % 28),
            data_final=date(2026, 12, 31),
            prioridade=1 + i % 4,
        )
        for i in range(n)
    ]


def build_model(
    path: Path,
    records: int = 5,
    table_rows: int | None = 20,
    *,
    with_table: bool = True,
    freeze: str | None = None,
    sheet_name: str = "Planilha1",
    headers: list[str] | None = None,
    formula_in_data: bool = False,
) -> Path:
    """Cria um modelo no padrão do real. ``table_rows`` = linhas de dados dentro da tabela."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    for col, header in zip("ABCDEF", headers or HEADERS):
        cell = ws[f"{col}1"]
        cell.value = header
        cell.border = BOX
        cell.font = Font(name="Aptos Narrow", size=11)
    for col, width in WIDTHS.items():
        ws.column_dimensions[col].width = width

    total_rows = max(table_rows or records, records)
    for i in range(total_rows):
        r = 2 + i
        for col in "ABCDEF":
            c = ws[f"{col}{r}"]
            c.border = BOX
            c.font = Font(name="Aptos Narrow", size=11)
            c.alignment = Alignment(horizontal="left" if col in "BC" else "center")
            if col in "DE":
                c.number_format = "mm-dd-yy"
        if i < records:
            ws[f"A{r}"] = 9000 + i
            ws[f"B{r}"] = f"Antiga {i}"
            ws[f"C{r}"] = f"Detalhe antigo {i}"
            ws[f"D{r}"] = datetime(2025, 1, 1 + i)
            ws[f"E{r}"] = datetime(2025, 12, 31)
            ws[f"F{r}"] = 1
    if formula_in_data:
        ws["A2"] = "=1+1"

    if freeze:
        ws.freeze_panes = freeze
    if with_table:
        table = Table(displayName="Tabela1", ref=f"A1:F{1 + total_rows}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium16", showRowStripes=True)
        ws.add_table(table)
    wb.save(path)
    return path


@pytest.fixture
def model_factory(tmp_path):
    counter = {"n": 0}

    def factory(**kwargs) -> Path:
        counter["n"] += 1
        return build_model(tmp_path / f"modelo_{counter['n']}.xlsx", **kwargs)

    return factory


@pytest.fixture
def model_path(model_factory) -> Path:
    return model_factory(records=5, table_rows=20)


BUNDLED_MODEL = Path(__file__).resolve().parent.parent / "assets" / "PROMOCOES-ATIVAS-modelo.xlsx"


@pytest.fixture
def real_model_path() -> Path:
    """Modelo real embutido no projeto (somente leitura; os testes não o modificam)."""
    assert BUNDLED_MODEL.is_file(), "modelo interno ausente em assets/"
    return BUNDLED_MODEL
