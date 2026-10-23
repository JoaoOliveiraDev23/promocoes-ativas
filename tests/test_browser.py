"""Teste OPCIONAL em navegador real: cola, clica uma vez e confere o download.

Só roda se o Playwright e um navegador estiverem instalados
(``pip install playwright`` e ``python -m playwright install chromium``);
caso contrário é ignorado automaticamente.
"""

from __future__ import annotations

import socket
import subprocess
import sys
import time
import urllib.request
from pathlib import Path

import pytest
from openpyxl import load_workbook

sync_api = pytest.importorskip("playwright.sync_api")
ROOT = Path(__file__).resolve().parent.parent


def _porta_livre() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


@pytest.fixture(scope="module")
def url():
    porta = _porta_livre()
    proc = subprocess.Popen(
        [sys.executable, "-m", "streamlit", "run", "app.py", "--server.headless", "true", "--server.port", str(porta)],
        cwd=ROOT,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    try:
        for _ in range(60):
            try:
                urllib.request.urlopen(f"http://127.0.0.1:{porta}/_stcore/health", timeout=1)
                break
            except OSError:
                time.sleep(0.5)
        else:
            pytest.skip("Servidor Streamlit não iniciou a tempo.")
        yield f"http://127.0.0.1:{porta}"
    finally:
        proc.terminate()
        proc.wait(timeout=10)


@pytest.fixture
def pagina(url):
    with sync_api.sync_playwright() as pw:
        try:
            navegador = pw.chromium.launch()
        except Exception:  # navegador não instalado
            pytest.skip("Chromium do Playwright não está instalado.")
        contexto = navegador.new_context(accept_downloads=True)
        page = contexto.new_page()
        page.goto(url)
        page.wait_for_selector("textarea")
        yield page
        navegador.close()


def test_um_clique_baixa_a_planilha_correta(pagina, tmp_path):
    texto = "\n".join(f"{3000 + i}\tPromoção {i}\tDetalhe {i}\t2026-02-{1 + i % 28:02d}\t2026-12-31\t1" for i in range(120))
    pagina.fill("textarea", texto)  # sem sair do campo antes de clicar
    with pagina.expect_download(timeout=20000) as info:
        pagina.get_by_role("button", name="Gerar e baixar planilha").click()
    download = info.value
    assert download.suggested_filename.startswith("PROMOCOES-ATIVAS ")
    destino = tmp_path / "saida.xlsx"
    download.save_as(destino)
    ws = load_workbook(destino).active
    assert ws["A2"].value == 3000 and ws.cell(121, 1).value == 3119 and ws.cell(122, 1).value is None
    assert ws.tables["Tabela1"].ref == "A1:F121"
